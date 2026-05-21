from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    warnings: list[str] = field(default_factory=list)


def _extract_plain(path: Path) -> ExtractionResult:
    return ExtractionResult(text=path.read_text(encoding="utf-8"))


def _extract_pdf(path: Path) -> ExtractionResult:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks: list[str] = []
    warnings: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
      text = page.extract_text() or ""
      if text.strip():
          chunks.append(text.strip())
      else:
          warnings.append(f"page {index} has no accessible text")
    return ExtractionResult(text="\n\n".join(chunks), warnings=warnings)


def _extract_docx(path: Path) -> ExtractionResult:
    from docx import Document

    document = Document(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    warnings: list[str] = []
    if document.tables:
        warnings.append("tables were skipped")
    return ExtractionResult(text="\n\n".join(paragraphs), warnings=warnings)


def extract_text(path: str | Path) -> ExtractionResult:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _extract_plain(source)
    if suffix == ".pdf":
        return _extract_pdf(source)
    if suffix == ".docx":
        return _extract_docx(source)
    raise ValueError(f"unsupported file type: {suffix or '<none>'}")
