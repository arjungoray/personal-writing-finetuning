from pathlib import Path

from extraction import extract_text


def test_extracts_txt_and_markdown(tmp_path: Path):
    text_file = tmp_path / "sample.txt"
    text_file.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")
    assert extract_text(text_file).text == "First paragraph.\n\nSecond paragraph."

    md_file = tmp_path / "sample.md"
    md_file.write_text("# Heading\n\nBody text.", encoding="utf-8")
    assert extract_text(md_file).text == "# Heading\n\nBody text."


def test_rejects_unsupported_file_type(tmp_path: Path):
    source = tmp_path / "sample.rtf"
    source.write_text("hello", encoding="utf-8")
    try:
        extract_text(source)
    except ValueError as error:
        assert "unsupported file type" in str(error)
    else:
        raise AssertionError("unsupported file type should fail")


def test_extracts_accessible_pdf_text(tmp_path: Path):
    from reportlab.pdfgen import canvas

    source = tmp_path / "sample.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.drawString(72, 720, "Accessible PDF sentence.")
    pdf.save()

    result = extract_text(source)
    assert "Accessible PDF sentence." in result.text
    assert result.warnings == []


def test_extracts_docx_paragraphs_and_warns_on_tables(tmp_path: Path):
    from docx import Document

    source = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Skipped table"
    document.save(source)

    result = extract_text(source)
    assert result.text == "First paragraph.\n\nSecond paragraph."
    assert result.warnings == ["tables were skipped"]
